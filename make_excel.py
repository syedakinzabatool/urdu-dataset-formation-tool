"""Makes output/dataset.xlsx from output/dataset.csv with text wrapping (a CSV cannot store wrap).
Run:  python make_excel.py      (close dataset.xlsx in Excel first)"""
import csv, math
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

out = Path(__file__).resolve().parent / "output"
rows = list(csv.reader(open(out / "dataset.csv", encoding="utf-8-sig", newline="")))
wb = Workbook(); ws = wb.active; ws.title = "dataset"
for r in rows:
    ws.append(r)
for col, width in {"A": 16, "B": 70, "C": 70}.items():
    ws.column_dimensions[col].width = width
for row in ws.iter_rows():
    for c in row:
        if c.column_letter == "B":   # Urdu: right-to-left, wrapped
            c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="right", readingOrder=2)
        else:
            c.alignment = Alignment(wrap_text=(c.column_letter == "C"), vertical="top")
for c in ws[1]:
    c.font = Font(bold=True)
ws.freeze_panes = "A2"
for i, r in enumerate(rows[1:], 2):
    n = max(math.ceil(len(r[1]) / 60), math.ceil(len(r[2]) / 62) if len(r) > 2 else 1, 1)
    ws.row_dimensions[i].height = 15 * n
wb.save(out / "dataset.xlsx")
print("saved", out / "dataset.xlsx", f"({len(rows) - 1} rows)")
